"""Fill a process checklist (.xlsx) from a markdown audit table.

The audit markdown must contain a table whose first column is the 1-based
row number in the checklist sheet, with "Status" and "Notes" columns:

    | Row | Item | Status | Notes |
    |---|---|---|---|
    | 2 | Business Purpose & Drivers | Done | ... |

Usage:
    python tools/checklist_fill.py <template.xlsx> <audit.md> <output.xlsx>

The template is never modified; the filled copy is written to <output.xlsx>.
"""

import re
import sys
import warnings

import openpyxl

ROW_RE = re.compile(r"^\|\s*(\d+)\s*\|")


def parse_audit(path):
    header = None
    rows = {}
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if header is None and "Status" in cells and "Notes" in cells:
                header = cells
                continue
            if header and not line.strip():
                if rows:
                    break  # only the first audit table counts
                continue
            if header and ROW_RE.match(line):
                rec = dict(zip(header, cells))
                rows[int(cells[0])] = (rec["Status"], rec["Notes"])
    if not rows:
        raise SystemExit(f"no audit rows found in {path}")
    return rows


def find_columns(ws):
    for row in ws.iter_rows(min_row=1, max_row=5):
        names = {str(c.value).strip().lower(): c.column for c in row if c.value}
        status = next((v for k, v in names.items() if k.startswith("status")), None)
        notes = next((v for k, v in names.items() if k.startswith("notes")), None)
        if status and notes:
            return status, notes
    raise SystemExit("could not find Status/Notes header in checklist")


def main(template, audit, output):
    warnings.filterwarnings("ignore")
    rows = parse_audit(audit)
    wb = openpyxl.load_workbook(template)
    ws = wb.worksheets[0]
    status_col, notes_col = find_columns(ws)
    for r, (status, notes) in rows.items():
        ws.cell(row=r, column=status_col, value=status)
        ws.cell(row=r, column=notes_col, value=notes)
    wb.save(output)
    print(f"filled {len(rows)} rows -> {output}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    main(*sys.argv[1:])
